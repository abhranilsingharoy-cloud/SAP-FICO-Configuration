import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_title_page(doc, title_text, subtitle_text):
    doc.add_paragraph('\n' * 5)
    title = doc.add_heading(title_text, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(24)
        run.bold = True
        
    doc.add_paragraph('\n' * 2)
    subtitle = doc.add_paragraph(subtitle_text)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in subtitle.runs:
        run.font.name = 'Arial'
        run.font.size = Pt(16)
        run.italic = True
    
    doc.add_page_break()

def add_heading(doc, text, level):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Arial'
        run.bold = True
        if level == 1:
            run.font.size = Pt(16)
        elif level == 2:
            run.font.size = Pt(14)
        else:
            run.font.size = Pt(12)

def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    if bold:
        run.bold = True
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(11)
    return p

doc = docx.Document()

# Title Page
add_title_page(doc, 'SAP FICO Configuration', 'A Comprehensive Guide for Beginners\n\nWeek 2 Task Submission')

# Introduction
add_heading(doc, '1. Introduction to SAP FICO', level=1)
add_paragraph(doc, "SAP FICO (Financial Accounting and Controlling) is one of the foundational modules in the SAP ERP system. It is designed to record all financial transactions that are posted by an entity and produce financial statements that are accurate at the end of the trading period. ")
add_paragraph(doc, "The configuration of SAP FICO is a critical step in any SAP implementation, as it lays down the organizational structure and accounting principles the company will follow. This guide outlines the essential step-by-step configurations needed to set up both the Financial Accounting (FI) and Controlling (CO) sub-modules.")

# FI Configuration
add_heading(doc, '2. Financial Accounting (FI) Configuration Steps', level=1)
add_paragraph(doc, "Financial Accounting focuses on external reporting. The configuration here defines the legal and organizational entities.")

add_heading(doc, '2.1 Define Enterprise Structure', level=2)
add_paragraph(doc, "The enterprise structure is the framework in which all business transactions are processed.")
add_bullet(doc, "Define Company (Transaction OX15): The company is the highest organizational unit in SAP. It groups together various company codes for consolidated financial reporting.")
add_bullet(doc, "Define Company Code (Transaction OX02): The company code represents an independent accounting entity (like a branch or a subsidiary). A legally required balance sheet and profit & loss statement are created at this level.")
add_bullet(doc, "Assign Company Code to Company (Transaction OX16): This step logically links the independent accounting entity to the larger corporate group.")

add_heading(doc, '2.2 Global Settings and Fiscal Year', level=2)
add_paragraph(doc, "These settings control how the financial year and posting periods are managed.")
add_bullet(doc, "Maintain Fiscal Year Variant (Transaction OB29): This defines the financial year (e.g., April to March, or Jan to Dec) and specifies the number of normal and special posting periods.")
add_bullet(doc, "Assign Company Code to Fiscal Year Variant (Transaction OB37): Links your specific company code to the chosen fiscal year structure.")
add_bullet(doc, "Define Variants for Open Posting Periods (Transaction OBBO): Used to control which specific accounting months are open for entering transactions.")
add_bullet(doc, "Open and Close Posting Periods (Transaction OB52): A critical monthly activity to ensure no back-dated postings occur in closed financial months.")

add_heading(doc, '2.3 Chart of Accounts and General Ledger', level=2)
add_paragraph(doc, "The Chart of Accounts (CoA) is the lifeblood of financial reporting, containing all General Ledger (GL) accounts.")
add_bullet(doc, "Define Chart of Accounts (Transaction OB13): Creates the master list of all GL accounts.")
add_bullet(doc, "Assign Company Code to Chart of Accounts (Transaction OB62): Ensures the company code uses the correct GL account list.")
add_bullet(doc, "Define Account Groups (Transaction OBD4): Groups similar accounts together (e.g., Fixed Assets, Current Liabilities) and controls the number ranges for these accounts.")
add_bullet(doc, "Define Retained Earnings Account (Transaction OB53): SAP requires a retained earnings account to automatically carry forward the balance of P&L accounts at year-end.")

add_heading(doc, '2.4 Document Control', level=2)
add_bullet(doc, "Define Document Types (Transaction OBA7): Differentiates between different accounting transactions (e.g., 'SA' for GL document, 'KR' for Vendor Invoice).")
add_bullet(doc, "Define Document Number Ranges (Transaction FBN1): Assigns a unique, sequential number to every financial posting to maintain a solid audit trail.")

# CO Configuration
add_heading(doc, '3. Controlling (CO) Configuration Steps', level=1)
add_paragraph(doc, "Controlling is focused on internal reporting. It provides management with the necessary data to analyze costs, revenues, and profitability.")

add_heading(doc, '3.1 Maintain Controlling Area', level=2)
add_bullet(doc, "Maintain Controlling Area (Transaction OKKP): The controlling area is the highest organizational unit within CO. You define the currency type, fiscal year variant, and activate specific sub-components like Cost Center Accounting or Profitability Analysis here.")
add_bullet(doc, "Assign Company Code to Controlling Area (Transaction OX19): You can assign one or more company codes to a single controlling area, provided they share the same Chart of Accounts and Fiscal Year Variant. This allows for cross-company cost reporting.")

add_heading(doc, '3.2 Cost Element Accounting', level=2)
add_bullet(doc, "Create Primary Cost Elements (Transaction KA01): Primary cost elements correspond directly to P&L expense accounts in FI. They act as the bridge transferring costs from FI to CO.")
add_bullet(doc, "Create Secondary Cost Elements (Transaction KA06): These do not have a corresponding FI account. They are used exclusively within CO to allocate costs between different cost centers (e.g., an IT department allocating server costs to other departments).")

add_heading(doc, '3.3 Cost Center Accounting', level=2)
add_bullet(doc, "Maintain Standard Hierarchy (Transaction OKEON): This is a tree-like structure grouping all cost centers in the controlling area. It is mandatory for cost center accounting.")
add_bullet(doc, "Create Cost Centers (Transaction KS01): Cost centers represent specific functional areas, departments, or locations that incur costs (e.g., HR Department, Marketing, Factory Unit 1).")
add_bullet(doc, "Maintain Number Ranges for Controlling Documents (Transaction KANK): Similar to FI, internal CO transactions (like assessments or distributions) require their own document numbers.")

# Conclusion
add_heading(doc, '4. Conclusion', level=1)
add_paragraph(doc, "The configuration of SAP FICO is a highly sequential and logical process. Starting with the broad enterprise structure, the configuration narrows down to specific rules governing accounts, periods, and cost centers. A precise and well-documented configuration at this stage ensures that all subsequent business transactions flow correctly into both the legal financial statements and internal management reports. This guide provides a robust baseline, preparing the system for more advanced integrations like Materials Management (MM) or Sales and Distribution (SD).")

doc.save('Detailed_SAP_FICO_Configuration_Report.docx')
